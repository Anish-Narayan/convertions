"""Read/write the mapping workbook (openpyxl only, to keep the .exe small).

Sheets
  Tables  : A1/B1 = pod labels (source pod, target pod); rows below = table names.
  Columns : source_table, target_table, source_column, target_column, status, match_method, confidence, notes
  Schemas : (optional) A1/B1 = pod labels; rows = schema placeholder in pod A -> pod B
Only rows whose status is 'confirmed' are applied by the converter (missing status column = confirmed).
"""
from dataclasses import dataclass
import openpyxl

TABLE_SHEET, COLUMN_SHEET, SCHEMA_SHEET = "Tables", "Columns", "Schemas"
COLUMN_HEADERS = ["source_table", "target_table", "source_column", "target_column",
                  "status", "match_method", "confidence", "notes"]


def _c(v):
    return "" if v is None else str(v).strip()


@dataclass
class Directed:
    table_rows: list
    column_rows: list
    schema_rows: list


class MappingBook:
    def __init__(self, labels, table_rows, column_rows, schema_rows, path=""):
        self.labels, self.table_rows = tuple(labels), table_rows
        self.column_rows, self.schema_rows, self.path = column_rows, schema_rows, path

    @classmethod
    def load(cls, path):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            names = wb.sheetnames
            ws = wb[TABLE_SHEET] if TABLE_SHEET in names else wb[names[0]]
            rows = [tuple((list(map(_c, r)) + ["", ""])[:2]) for r in ws.iter_rows(values_only=True)]
            if not rows or not rows[0][0] or not rows[0][1]:
                raise ValueError("First row of the table sheet must hold the two pod labels in columns A and B.")
            labels = rows[0]
            if labels[0].lower() == labels[1].lower():
                raise ValueError("The two pod labels in row 1 must be different.")
            tables = [r for r in rows[1:] if r[0] and r[1]]
            cols, schemas = [], []
            if COLUMN_SHEET in names:
                it = wb[COLUMN_SHEET].iter_rows(values_only=True)
                head = [_c(h).lower() for h in next(it, [])]
                for r in it:
                    d = {h: _c(v) for h, v in zip(head, r) if h}
                    if d.get("source_table") and d.get("target_table"):
                        d.setdefault("status", "confirmed")
                        d["status"] = d["status"] or "confirmed"
                        cols.append(d)
            if SCHEMA_SHEET in names:
                srows = [tuple((list(map(_c, r)) + ["", ""])[:2]) for r in wb[SCHEMA_SHEET].iter_rows(values_only=True)]
                schemas = [r for r in srows[1:] if r[0] and r[1]]
            return cls(labels, tables, cols, schemas, str(path))
        finally:
            wb.close()

    def directed(self, source, target):
        lo = [x.lower() for x in self.labels]
        s, t = source.lower(), target.lower()
        if s == t:
            raise ValueError("Source and target pod must be different.")
        if (s, t) == (lo[0], lo[1]):
            return Directed(list(self.table_rows), list(self.column_rows), list(self.schema_rows))
        if (s, t) == (lo[1], lo[0]):
            swap = {"source_table": "target_table", "target_table": "source_table",
                    "source_column": "target_column", "target_column": "source_column"}
            cols = [{swap.get(k, k): v for k, v in r.items()} for r in self.column_rows]
            return Directed([(b, a) for a, b in self.table_rows], cols, [(b, a) for a, b in self.schema_rows])
        raise ValueError(f"Selection {source!r} -> {target!r} is not covered by this workbook.")


def write_mapping_workbook(path, labels, table_rows, column_rows, schema_rows=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = TABLE_SHEET
    ws.append(list(labels))
    for r in table_rows:
        ws.append(list(r))
    ws = wb.create_sheet(COLUMN_SHEET)
    ws.append(COLUMN_HEADERS)
    for r in column_rows:
        ws.append([r.get(h, "") for h in COLUMN_HEADERS])
    ws = wb.create_sheet(SCHEMA_SHEET)
    ws.append(list(labels))
    for r in schema_rows:
        ws.append(list(r))
    wb.save(path)


def write_report_workbook(path, sheets):
    """sheets: {name: (headers, rows)}"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.append(headers)
        for r in rows:
            ws.append(["" if v is None else v for v in r])
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(14, min(50, len(str(h)) + 6))
    wb.save(path)
